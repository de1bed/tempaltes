"""
Convierte las plantillas de corridas (src/corridas/plantillas/*.xlsx) a JSON
para el motor de cálculo de la app: por hoja, cada celda con su fórmula (tal
cual está en el Excel aprobado) y su último valor calculado.

Uso (cuando Nóminas mande una plantilla actualizada, reemplazar el .xlsx y correr):
    python3 scripts/extraer-plantillas-corridas.py

Requiere openpyxl (pip install openpyxl). openpyxl expande las fórmulas
compartidas, así cada celda queda con su propia fórmula.
"""

import datetime
import json
import pathlib

import openpyxl

CARPETA = pathlib.Path(__file__).resolve().parent.parent / "src" / "corridas" / "plantillas"
ERRORES = {"#N/A", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#NULL!"}
EPOCA = datetime.datetime(1899, 12, 30)


def valor(v):
    """Valor guardado en el Excel → JSON. Fechas como número de serie de Excel."""
    if isinstance(v, datetime.datetime):
        return {"n": (v - EPOCA).total_seconds() / 86400}
    if isinstance(v, datetime.date):
        return {"n": (datetime.datetime(v.year, v.month, v.day) - EPOCA).days}
    if isinstance(v, bool):
        return {"b": v}
    if isinstance(v, (int, float)):
        return {"n": v}
    if isinstance(v, str):
        return {"e": v} if v in ERRORES else {"s": v}
    return None


def extraer(xlsx: pathlib.Path) -> dict:
    formulas = openpyxl.load_workbook(xlsx)
    valores = openpyxl.load_workbook(xlsx, data_only=True)
    hojas = {}
    for ws in formulas:
        celdas = {}
        for fila in ws.iter_rows():
            for c in fila:
                v = valores[ws.title][c.coordinate].value
                if isinstance(c.value, str) and c.value.startswith("="):
                    celda = {"f": c.value[1:]}
                    cache = valor(v)
                    if cache is not None:
                        celda["c"] = cache
                    celdas[c.coordinate] = celda
                elif c.value is not None:
                    celdas[c.coordinate] = valor(c.value)
        hojas[ws.title] = celdas
    return {"archivo": xlsx.name, "hojas": hojas}


if __name__ == "__main__":
    for xlsx in sorted(CARPETA.glob("*.xlsx")):
        destino = xlsx.with_suffix(".json")
        destino.write_text(json.dumps(extraer(xlsx), ensure_ascii=False, separators=(",", ":")))
        print(f"{xlsx.name} → {destino.name} ({destino.stat().st_size // 1024} KB)")
